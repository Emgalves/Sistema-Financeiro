"""
Qualificação do CONTRATANTE (cliente) nos contratos de empreitada e de
administração de obra.

Regra definida em 30/09/2026:
- Cliente PJ (CNPJ na coluna G de Clientes.xlsx): sem estado civil e sem
  residência. O endereço da SEDE vem da coluna ENDEREÇO da base de
  fornecedores, a mesma fonte já usada para a contratada. Se o cliente PJ
  não estiver na base de fornecedores, ou estiver lá sem endereço, o
  contrato NÃO é gerado e o usuário recebe uma mensagem dizendo o que fazer.
- Cliente PF: como antes (CPF, estado civil, residência). A residência é a
  coluna AA (DOMICILIO) de Clientes.xlsx, se preenchida; senão, o endereço
  da obra (coluna B), como sempre foi.

Usado por:
- src/modules/gerador_contrato.py                 (empreitada)
- src/gestao_medicoes.py                          (tela que gera a empreitada)
- src/taxas_administracao/gerador_contrato_adm.py (administração de obra)
"""
from openpyxl import load_workbook


class ContratantePJSemCadastro(Exception):
    """Cliente PJ sem cadastro (ou sem endereço) na base de fornecedores."""


def _doc_limpo(valor):
    """Letras e dígitos do documento, em maiúsculas. Trata célula numérica."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return ''.join(c for c in str(valor or '').upper() if c.isalnum())


def eh_pj(documento):
    """True somente se o documento tiver o formato de um CNPJ:
    - numérico com 12 a 14 dígitos (mesma regra do cadastro de clientes); ou
    - alfanumérico com 14 posições e os 2 dígitos verificadores numéricos.
    Textos de preenchimento como "não informado" ou "nan" (célula vazia)
    NÃO contam como CNPJ: o cliente é tratado como PF, como antes."""
    limpo = _doc_limpo(documento)
    if limpo.isdigit():
        return 11 < len(limpo) <= 14
    return len(limpo) == 14 and limpo[12:].isdigit() and any(c.isdigit() for c in limpo[:12])


def buscar_sede_pj(nome_cliente, cnpj, arquivo_fornecedores):
    """Retorna o ENDEREÇO (sede) do CNPJ na base de fornecedores.

    Lança ContratantePJSemCadastro, com a mensagem pronta para o usuário,
    se o CNPJ não estiver na base ou estiver sem endereço.
    """
    alvo = _doc_limpo(cnpj).zfill(14)

    wb = load_workbook(arquivo_fornecedores, read_only=True, data_only=True)
    try:
        ws = wb['Fornecedores']
        cabecalho = [str(c or '').strip().upper()
                     for c in next(ws.iter_rows(max_row=1, values_only=True))]
        col_end = next((i for i, c in enumerate(cabecalho)
                        if c in ('ENDEREÇO', 'ENDERECO')), 15)  # padrão: coluna P
        col_nome = next((i for i, c in enumerate(cabecalho) if c == 'NOME'), 3)

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            if _doc_limpo(row[0]).zfill(14) != alvo:
                continue
            endereco = str(row[col_end]).strip() if len(row) > col_end and row[col_end] else ''
            if endereco:
                return endereco
            nome_forn = str(row[col_nome]).strip() if len(row) > col_nome and row[col_nome] else ''
            raise ContratantePJSemCadastro(
                f"O cliente '{nome_cliente}' é pessoa jurídica (CNPJ {cnpj}) e está "
                f"na base de fornecedores{f' como {nome_forn}' if nome_forn else ''}, "
                f"mas sem endereço.\n\n"
                f"O endereço da sede usado no contrato vem de lá. Preencha o campo "
                f"Endereço em Fornecedor → Editar Fornecedor e gere o contrato novamente."
            )
    finally:
        wb.close()

    raise ContratantePJSemCadastro(
        f"O cliente '{nome_cliente}' é pessoa jurídica (CNPJ {cnpj}) e não está "
        f"cadastrado na base de fornecedores.\n\n"
        f"O endereço da sede usado no contrato vem de lá. Cadastre a empresa em "
        f"Fornecedor → Novo Fornecedor, com o endereço da sede, e gere o "
        f"contrato novamente."
    )
