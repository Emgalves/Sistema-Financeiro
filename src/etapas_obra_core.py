# -*- coding: utf-8 -*-
"""
etapas_obra_core.py
===================
Núcleo do "Relatório por Etapa da Obra" e da ferramenta de regularização de etapas.

Este módulo NÃO usa tkinter nem reportlab: contém somente regras de negócio, leitura da
planilha do cliente, agregações, sugestões e gravação segura. Assim ele pode ser testado
isoladamente (ver tests/teste_etapas_obra.py).

Regras de classificação de cada lançamento (aba Dados, STATUS diferente de EXCLUIDO):

    Categoria TAX, ADM ou TP (qualquer data, mesmo com etapa escrita)
        -> CUSTOS GERAIS, linha 1 ("regra")
    ETAPA_OBRA = "CUSTOS GERAIS DA OBRA"
        -> CUSTOS GERAIS, linha 2 ("marcados")
    ETAPA_OBRA = etapa da lista de Configurações
        -> essa etapa
    ETAPA_OBRA vazia e DATA_REL >= marco do cliente (ou etapa escrita fora da lista)
        -> CUSTOS GERAIS, linha 3 ("sem etapa informada": pendência)
    ETAPA_OBRA vazia e DATA_REL < marco do cliente
        -> "Anteriores ao controle por etapa" (permanecem como estão)

O marco é uma data por cliente, escolhida entre as quinzenas do próprio cliente (aba
ETAPAS_OBRA). Enquanto o cliente não tiver marco salvo, vale MARCO_PADRAO.

Gravação: somente ETAPA_OBRA (coluna Q) e HISTORICO_ALTERACAO (coluna P) de linhas com Q vazia,
mais a aba ETAPAS_OBRA (ordem e marco). Sempre com backup, conferência prévia e troca atômica.
"""

import logging
import os
import re
import shutil
import unicodedata
import warnings
from dataclasses import dataclass, field
from datetime import datetime, date
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook

logger = logging.getLogger(__name__)
warnings.filterwarnings("ignore", message=".*image format is not supported.*")

# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------
ETAPA_GERAIS = "CUSTOS GERAIS DA OBRA"
CATS_REGRA = ("TAX", "ADM", "TP")
CATS_ETAPA = ("MO", "MAT", "LOC", "SERV", "DIV")
CATS_TODAS = ("MO", "MAT", "LOC", "SERV", "DIV", "ADM", "TAX", "TP")
NOMES_CATEGORIA = {
    "MO": "Mão de obra", "MAT": "Material", "LOC": "Locação", "SERV": "Serviços",
    "DIV": "Diversos", "ADM": "Administração", "TAX": "Taxa de administração", "TP": "Tarifas públicas",
}
SUBGRUPOS_MO = ("REMUNERAÇÃO", "BENEFÍCIOS", "ENCARGOS", "DIVERSOS", "OUTROS")

# Data padrão do marco enquanto o cliente não tiver um marco salvo (entrada em produção do campo
# ETAPA_OBRA). Cada cliente pode recuar (ou avançar) o marco pela tela "Ordem e marco".
MARCO_PADRAO = date(2025, 9, 5)

ABA_DADOS = "Dados"
ABA_ETAPAS = "ETAPAS_OBRA"
PASTA_BACKUP = "Backups_Sistema"
LIMITE_DOMINANCIA = 0.80

# Posições fixas das colunas gravadas pelo sistema (O, P, Q)
COL_ID, COL_HISTORICO, COL_ETAPA = 15, 16, 17

# Lista de reserva (usada só se GerenciadorConfiguracoes não puder ser importado)
ETAPAS_RESERVA = [
    "ACABAMENTOS", "ALVENARIA", "AUTOMAÇÃO", "CHAPISCO E REBOCO", "COBERTURA", "CONTRAPISO", "DEMOLIÇÃO",
    "EMBOÇO E REBOCO", "ESQUADRIAS", "ESTRUTURA", "FINALIZANDO", "IMPERMEABILIZAÇÃO", "INFRA DE GÁS",
    "INFRAESTRUTURA", "INSTALAÇÃO DA OBRA", "INSTALAÇÕES ELÉTRICAS", "INSTALAÇÕES HIDRÁULICAS",
    "LIMPEZA FINAL", "MURO DE ARRIMO", "PAISAGISMO", "PINTURA", "PISCINA", "PISOS", "REBOCO", "REFORMA",
    "REVESTIMENTOS", "SERVIÇOS PRELIMINARES", "SUPRAESTRUTURA", "TERRAPLANAGEM", "VEDAÇÃO", ETAPA_GERAIS,
]

# Destinos (primeiro nível) de cada lançamento
D_ETAPA = "ETAPA"
D_GERAIS_REGRA = "GERAIS_REGRA"
D_GERAIS_MARCADO = "GERAIS_MARCADO"
D_SEM_ETAPA = "SEM_ETAPA"
D_ANTERIOR = "ANTERIOR"

_COLUNAS_PADRAO = {
    "DATA_REL": 1, "TP_DESP": 2, "CNPJ_CPF": 3, "NOME": 4, "REFERENCIA": 5, "NF": 6, "VR_UNIT": 7,
    "DIAS": 8, "VALOR": 9, "DT_VENCTO": 10, "CATEGORIA": 11, "DADOS_BANCARIOS": 12, "OBSERVACAO": 13,
    "STATUS": 14, "ID_LANCAMENTO": 15, "HISTORICO_ALTERACAO": 16, "ETAPA_OBRA": 17, "INSUMO": 18,
}


class ArquivoEmUsoError(Exception):
    """A planilha está aberta em outro programa (ou bloqueada) e não pode ser gravada."""


class PlanilhaAlteradaError(Exception):
    """A planilha foi alterada por outro usuário durante a gravação; nada foi gravado."""


# ---------------------------------------------------------------------------
# Utilitários
# ---------------------------------------------------------------------------
def normalizar(txt) -> str:
    """Maiúsculas, sem acentos e sem espaços duplicados (para comparar textos)."""
    if txt is None:
        return ""
    s = unicodedata.normalize("NFKD", str(txt))
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", s).strip().upper()


def fmt_moeda(v) -> str:
    try:
        return f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "0,00"


def fmt_pct(v) -> str:
    try:
        return f"{float(v) * 100:,.1f}%".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "0,0%"


def fmt_data(d) -> str:
    if d is None or (isinstance(d, float) and pd.isna(d)) or pd.isna(d):
        return ""
    return pd.Timestamp(d).strftime("%d/%m/%Y")


def parse_data(txt) -> Optional[pd.Timestamp]:
    """Aceita dd/mm/aaaa, aaaa-mm-dd, date ou datetime. Devolve None se inválida."""
    if txt is None or str(txt).strip() == "":
        return None
    if isinstance(txt, (date, datetime, pd.Timestamp)):
        return pd.Timestamp(txt).normalize()
    s = str(txt).strip()
    for f in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return pd.Timestamp(datetime.strptime(s, f))
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------------------
# Lista de etapas e contexto
# ---------------------------------------------------------------------------
def obter_lista_etapas() -> List[str]:
    """Lista de etapas cadastrada em Configurações (a mesma usada na entrada de dados)."""
    try:
        from src.configuracoes_sistema import GerenciadorConfiguracoes
        lista = [str(x).strip() for x in GerenciadorConfiguracoes.get_etapas_obra() if str(x).strip()]
        if lista:
            return lista
    except Exception as e:  # noqa: BLE001
        logger.warning(f"Lista de etapas: usando lista interna de reserva ({e})")
    return list(ETAPAS_RESERVA)


@dataclass
class Contexto:
    lista: List[str]
    marco: pd.Timestamp
    etapas_norm: Dict[str, str] = field(default_factory=dict)  # texto normalizado -> nome canônico
    gerais: str = ETAPA_GERAIS
    gerais_norm: str = ""
    gerais_cadastrada: bool = False

    def __post_init__(self):
        self.gerais_norm = normalizar(ETAPA_GERAIS)
        self.etapas_norm = {}
        for e in self.lista:
            n = normalizar(e)
            if n == self.gerais_norm:
                self.gerais = e          # usa o texto exatamente como foi cadastrado
                self.gerais_cadastrada = True
            else:
                self.etapas_norm[n] = e

    def etapas_para_combo(self) -> List[str]:
        """Etapas (ordem alfabética) e, ao final, CUSTOS GERAIS DA OBRA."""
        return sorted(self.etapas_norm.values(), key=normalizar) + [self.gerais]

    def canonica(self, nome) -> Optional[str]:
        """Nome canônico de uma etapa (ou de CUSTOS GERAIS); None se não reconhecida."""
        n = normalizar(nome)
        if n == self.gerais_norm:
            return self.gerais
        return self.etapas_norm.get(n)


def criar_contexto(arquivo=None, marco=None, lista=None) -> Contexto:
    """Contexto do cliente: lista de etapas + marco (salvo no cliente, informado ou padrão)."""
    if marco is None and arquivo is not None:
        try:
            marco = ler_config_cliente(arquivo).get("marco")
        except Exception as e:  # noqa: BLE001
            logger.warning(f"Não foi possível ler ETAPAS_OBRA: {e}")
    m = parse_data(marco) or pd.Timestamp(MARCO_PADRAO)
    return Contexto(lista=list(lista) if lista else obter_lista_etapas(), marco=m)


# ---------------------------------------------------------------------------
# Leitura da planilha
# ---------------------------------------------------------------------------
def _mapa_colunas(ws) -> Dict[str, int]:
    """Cabeçalho -> nº da coluna (normalizado), com as posições padrão como reserva."""
    mapa = {}
    for c in range(1, min(ws.max_column, 40) + 1):
        v = ws.cell(row=1, column=c).value
        if v not in (None, ""):
            mapa[normalizar(v)] = c
    cols = dict(_COLUNAS_PADRAO)
    for nome in cols:
        if nome in mapa:
            cols[nome] = mapa[nome]
    return cols


def carregar_lancamentos(arquivo) -> pd.DataFrame:
    """
    Lê a aba Dados. Uma linha por lançamento (inclui EXCLUIDO; filtre depois).
    LINHA = número da linha na planilha (Vinculacoes.Linha_Lancamento usa este número).
    """
    wb = load_workbook(arquivo, data_only=True)
    try:
        if ABA_DADOS not in wb.sheetnames:
            raise ValueError(f"A planilha não possui a aba '{ABA_DADOS}'.")
        ws = wb[ABA_DADOS]
        cols = _mapa_colunas(ws)
        idx = {n: c - 1 for n, c in cols.items()}
        registros = []
        for linha, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            def g(nome):
                i = idx[nome]
                return row[i] if i < len(row) else None
            valor, data = g("VALOR"), g("DATA_REL")
            if valor in (None, "") and data in (None, ""):
                continue
            registros.append({
                "LINHA": linha, "ID": g("ID_LANCAMENTO"), "DATA_REL": data, "TP_DESP": g("TP_DESP"),
                "NOME": g("NOME"), "REFERENCIA": g("REFERENCIA"), "VALOR": valor,
                "CATEGORIA": g("CATEGORIA"), "STATUS": g("STATUS"), "ETAPA_RAW": g("ETAPA_OBRA"),
                "HISTORICO": g("HISTORICO_ALTERACAO"),
            })
    finally:
        wb.close()
    df = pd.DataFrame(registros)
    if df.empty:
        return df
    df["DATA_REL"] = pd.to_datetime(df["DATA_REL"], errors="coerce", dayfirst=True)
    df["VALOR"] = pd.to_numeric(df["VALOR"], errors="coerce").fillna(0.0)
    df["ID"] = pd.to_numeric(df["ID"], errors="coerce")
    df["TP_DESP"] = pd.to_numeric(df["TP_DESP"], errors="coerce")
    for c in ("NOME", "REFERENCIA", "CATEGORIA", "STATUS", "ETAPA_RAW"):
        df[c] = df[c].fillna("").astype(str).str.strip()
    df["CATEGORIA"] = df["CATEGORIA"].str.upper()
    df["STATUS"] = df["STATUS"].str.upper()
    df["NOME"] = df["NOME"].str.replace(r"\s+", " ", regex=True)
    return df


def preparar(df: pd.DataFrame, ctx: Contexto, posicao=None) -> pd.DataFrame:
    """Exclui os EXCLUIDO, aplica a data de posição e roteia cada lançamento (coluna DESTINO)."""
    if df is None or df.empty:
        return pd.DataFrame(columns=["LINHA", "DESTINO", "ETAPA_FINAL", "ETAPA_INVALIDA", "VALOR", "CATEGORIA", "DATA_REL"])
    d = df[df["STATUS"] != "EXCLUIDO"].copy()
    pos = parse_data(posicao)
    if pos is not None:
        d = d[d["DATA_REL"] <= pos]
    destinos, finais, invalidas = [], [], []
    for cat, raw, data in zip(d["CATEGORIA"], d["ETAPA_RAW"], d["DATA_REL"]):
        dest, final, inval = _rotear(cat, raw, data, ctx)
        destinos.append(dest); finais.append(final); invalidas.append(inval)
    d["DESTINO"], d["ETAPA_FINAL"], d["ETAPA_INVALIDA"] = destinos, finais, invalidas
    return d.reset_index(drop=True)


def _rotear(cat, raw, data, ctx: Contexto) -> Tuple[str, str, bool]:
    if cat in CATS_REGRA:
        return D_GERAIS_REGRA, "", False
    if raw:
        n = normalizar(raw)
        if n == ctx.gerais_norm:
            return D_GERAIS_MARCADO, "", False
        if n in ctx.etapas_norm:
            return D_ETAPA, ctx.etapas_norm[n], False
        return D_SEM_ETAPA, "", True          # etapa escrita que não está na lista
    if pd.notna(data) and data < ctx.marco:
        return D_ANTERIOR, "", False
    return D_SEM_ETAPA, "", False


# ---------------------------------------------------------------------------
# Configuração por cliente (aba ETAPAS_OBRA): ordem das etapas e marco
# ---------------------------------------------------------------------------
def ler_config_cliente(arquivo) -> dict:
    """{'ordem': [etapas], 'marco': Timestamp|None}. Sem a aba, devolve vazio."""
    wb = load_workbook(arquivo, read_only=True, data_only=True)
    try:
        if ABA_ETAPAS not in wb.sheetnames:
            return {"ordem": [], "marco": None}
        ws = wb[ABA_ETAPAS]
        ordem, marco = [], None
        for row in ws.iter_rows(min_row=2, values_only=True):
            row = list(row) + [None] * 6
            if row[1] not in (None, ""):
                ordem.append((row[0] if isinstance(row[0], (int, float)) else 10 ** 6, str(row[1]).strip()))
            if normalizar(row[3]) == "MARCO_DATA" and row[4] not in (None, ""):
                marco = parse_data(row[4])
        ordem.sort(key=lambda t: t[0])
        return {"ordem": [e for _, e in ordem], "marco": marco}
    finally:
        wb.close()


def salvar_config_cliente(arquivo, ordem: List[str], marco) -> str:
    """Grava ordem e marco na aba ETAPAS_OBRA (criada se não existir). Devolve o caminho do backup."""
    m = parse_data(marco)

    def editar(wb):
        ws = wb[ABA_ETAPAS] if ABA_ETAPAS in wb.sheetnames else wb.create_sheet(ABA_ETAPAS)
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
        ws.cell(row=1, column=1, value="ORDEM"); ws.cell(row=1, column=2, value="ETAPA")
        ws.cell(row=1, column=4, value="PARAMETRO"); ws.cell(row=1, column=5, value="VALOR")
        for i, e in enumerate(ordem, start=1):
            ws.cell(row=i + 1, column=1, value=i); ws.cell(row=i + 1, column=2, value=e)
        ws.cell(row=2, column=4, value="MARCO_DATA")
        if m is not None:
            c = ws.cell(row=2, column=5, value=m.to_pydatetime()); c.number_format = "dd/mm/yyyy"
        ws.column_dimensions["B"].width = 32; ws.column_dimensions["D"].width = 16; ws.column_dimensions["E"].width = 14
        return None

    _, backup = _editar_planilha(arquivo, editar, "ANTES_ORDEM_MARCO")
    return backup


# ---------------------------------------------------------------------------
# Ordem das etapas
# ---------------------------------------------------------------------------
def ordem_cronologica(df_prep: pd.DataFrame) -> List[str]:
    """Etapas em uso, ordenadas pela data mediana dos lançamentos (desempate: primeira data)."""
    e = df_prep[df_prep["DESTINO"] == D_ETAPA]
    if e.empty:
        return []
    linhas = []
    for nome, g in e.groupby("ETAPA_FINAL"):
        s = g["DATA_REL"].dropna().sort_values()
        if s.empty:
            linhas.append((pd.Timestamp.max, pd.Timestamp.max, nome)); continue
        linhas.append((s.iloc[len(s) // 2], s.iloc[0], nome))
    linhas.sort(key=lambda t: (t[0], t[1], normalizar(t[2])))
    return [n for _, _, n in linhas]


def ordem_efetiva(df_prep: pd.DataFrame, ordem_salva: Optional[List[str]], ctx: Contexto) -> Tuple[List[str], str]:
    """
    Ordem final das etapas do relatório. Etapas da ordem salva primeiro (na ordem salva);
    etapas em uso que não estão na ordem salva entram ao final, em ordem cronológica.
    Devolve (ordem, origem) com origem 'salva' ou 'cronológica'.
    """
    cron = ordem_cronologica(df_prep)
    if not ordem_salva:
        return cron, "cronológica"
    salvas = [ctx.canonica(e) for e in ordem_salva]
    salvas = [e for e in salvas if e and e in cron]
    resto = [e for e in cron if e not in salvas]
    return salvas + resto, "salva"


# ---------------------------------------------------------------------------
# Subgrupos da mão de obra
# ---------------------------------------------------------------------------
_RE_ENCARGOS = re.compile(r"\bINSS\b|\bFGTS\b|\bIRRF\b")
_RE_DIVERSOS_NOME = re.compile(
    r"FOLHA|MENSALIDADE|\bMHS\b|\bSST\b|E-?SOCIAL|\bPASI\b|WORK ?MED|UNIFORME|\bEPIS?\b|MOTOBOY|FRETE")
_RE_BENEFICIOS = re.compile(r"TRANSPORTE|CAFE|CESTA|\bVT\b")
_RE_REMUNERACAO = re.compile(r"SALARIO|DIARIA|FERIAS|RESCIS|13")
_RE_DIVERSOS_REF = re.compile(r"UNIFORME|EQUIPAMENTO|EXAME|SEGURO|FRETE|BOTAS|\bCAPA\b|\bREF\b")


def subgrupo_mo(nome, referencia) -> str:
    """
    Subgrupo da mão de obra. O beneficiário (NOME) é avaliado antes da referência: assim
    "FOLHA DP" com referência "REF. 13º SALÁRIO" fica em DIVERSOS, e não em REMUNERAÇÃO.
    O que não se encaixa em nenhuma regra vai para OUTROS (rede de segurança: o total fecha sempre).
    """
    n, r = normalizar(nome), normalizar(referencia)
    if _RE_ENCARGOS.search(n):
        return "ENCARGOS"
    if _RE_DIVERSOS_NOME.search(n):
        return "DIVERSOS"
    if _RE_BENEFICIOS.search(r):
        return "BENEFÍCIOS"
    if _RE_REMUNERACAO.search(r):
        return "REMUNERAÇÃO"
    if _RE_DIVERSOS_REF.search(r):
        return "DIVERSOS"
    return "OUTROS"


# ---------------------------------------------------------------------------
# Dados do relatório
# ---------------------------------------------------------------------------
@dataclass
class BlocoEtapa:
    nome: str
    primeira: pd.Timestamp
    ultima: pd.Timestamp
    n: int
    total: float
    cats: Dict[str, float]
    mo_sub: Dict[str, float]
    fornecedores: List[Tuple[str, float]]


@dataclass
class BlocoSimples:
    n: int
    total: float
    cats: Dict[str, float]
    fornecedores: List[Tuple[str, float]]


@dataclass
class DadosRelatorio:
    arquivo: str
    posicao: pd.Timestamp
    inicio: pd.Timestamp
    total: float
    marco: pd.Timestamp
    ordem_origem: str
    etapas: List[BlocoEtapa]
    regra: BlocoSimples
    marcados: BlocoSimples
    sem_etapa: BlocoSimples
    anteriores: BlocoSimples
    n_lancamentos: int
    avisos: List[str] = field(default_factory=list)

    @property
    def total_gerais(self) -> float:
        return self.regra.total + self.marcados.total + self.sem_etapa.total

    def conferir(self):
        soma = sum(b.total for b in self.etapas) + self.regra.total + self.marcados.total \
            + self.sem_etapa.total + self.anteriores.total
        if abs(soma - self.total) > 0.005:
            raise ValueError(
                f"Conferência de totais falhou: soma dos blocos R$ {fmt_moeda(soma)} "
                f"difere do total R$ {fmt_moeda(self.total)}.")


def _por_categoria(x: pd.DataFrame, categorias) -> Dict[str, float]:
    """Soma por categoria; categorias desconhecidas entram em DIV para o total sempre fechar."""
    out = {c: 0.0 for c in categorias}
    for cat, v in x.groupby("CATEGORIA")["VALOR"].sum().items():
        out[cat if cat in out else ("DIV" if "DIV" in out else list(out)[-1])] += float(v)
    return out


def _top_fornecedores(x: pd.DataFrame, n=5) -> List[Tuple[str, float]]:
    f = x[x["CATEGORIA"].isin(["MAT", "LOC", "SERV", "DIV"])].groupby("NOME")["VALOR"].sum()
    return [(str(k), float(v)) for k, v in f.sort_values(ascending=False).head(n).items() if k]


def montar_relatorio(arquivo, posicao=None, ctx: Optional[Contexto] = None) -> DadosRelatorio:
    """Carrega a planilha, aplica as regras e devolve os dados do relatório (já conferidos)."""
    ctx = ctx or criar_contexto(arquivo)
    cfg = ler_config_cliente(arquivo)
    bruto = carregar_lancamentos(arquivo)
    if bruto.empty:
        raise ValueError("A aba Dados não possui lançamentos.")
    ativos = bruto[bruto["STATUS"] != "EXCLUIDO"]
    pos = parse_data(posicao) or ativos["DATA_REL"].max()
    d = preparar(bruto, ctx, pos)
    if d.empty:
        raise ValueError("Não há lançamentos ativos até a data de posição informada.")
    ordem, origem = ordem_efetiva(d, cfg.get("ordem"), ctx)

    etapas = []
    for nome in ordem:
        x = d[(d["DESTINO"] == D_ETAPA) & (d["ETAPA_FINAL"] == nome)]
        cats = _por_categoria(x, CATS_ETAPA)
        mo = x[x["CATEGORIA"] == "MO"]
        mo_sub = {s: 0.0 for s in SUBGRUPOS_MO}
        for nm, rf, v in zip(mo["NOME"], mo["REFERENCIA"], mo["VALOR"]):
            mo_sub[subgrupo_mo(nm, rf)] += float(v)
        etapas.append(BlocoEtapa(nome, x["DATA_REL"].min(), x["DATA_REL"].max(), len(x), float(x["VALOR"].sum()),
                                 cats, mo_sub, _top_fornecedores(x)))

    def simples(destino):
        x = d[d["DESTINO"] == destino]
        return BlocoSimples(len(x), float(x["VALOR"].sum()), _por_categoria(x, CATS_TODAS), _top_fornecedores(x))

    dados = DadosRelatorio(
        arquivo=str(arquivo), posicao=pd.Timestamp(pos), inicio=ativos["DATA_REL"].min(),
        total=float(d["VALOR"].sum()), marco=ctx.marco, ordem_origem=origem, etapas=etapas,
        regra=simples(D_GERAIS_REGRA), marcados=simples(D_GERAIS_MARCADO), sem_etapa=simples(D_SEM_ETAPA),
        anteriores=simples(D_ANTERIOR), n_lancamentos=len(d))
    if not ctx.gerais_cadastrada:
        dados.avisos.append(f"'{ETAPA_GERAIS}' não está cadastrada na lista de etapas de Configurações.")
    if int(d["ETAPA_INVALIDA"].sum()):
        dados.avisos.append(f"{int(d['ETAPA_INVALIDA'].sum())} lançamento(s) com etapa fora da lista foram "
                            f"tratados como 'sem etapa informada'.")
    dados.conferir()
    return dados


def resumo_situacao(arquivo, ctx: Optional[Contexto] = None) -> dict:
    """Números do painel: pendências, anteriores ao marco, etapas em uso."""
    ctx = ctx or criar_contexto(arquivo)
    d = preparar(carregar_lancamentos(arquivo), ctx)
    total = float(d["VALOR"].sum()) if len(d) else 0.0
    sem = d[d["DESTINO"] == D_SEM_ETAPA]
    ant = d[d["DESTINO"] == D_ANTERIOR]
    cfg = ler_config_cliente(arquivo)
    return {
        "total": total, "ultima_data": d["DATA_REL"].max() if len(d) else None, "marco": ctx.marco,
        "n_sem": len(sem), "v_sem": float(sem["VALOR"].sum()),
        "n_sem_mo": int((sem["CATEGORIA"] == "MO").sum()), "n_sem_outros": int((sem["CATEGORIA"] != "MO").sum()),
        "n_ant": len(ant), "v_ant": float(ant["VALOR"].sum()),
        "n_etapas": int(d.loc[d["DESTINO"] == D_ETAPA, "ETAPA_FINAL"].nunique()) if len(d) else 0,
        "ordem": "salva" if cfg.get("ordem") else "cronológica",
        "gerais_cadastrada": ctx.gerais_cadastrada,
    }


def quinzenas_do_cliente(df: pd.DataFrame) -> List[pd.Timestamp]:
    """Datas de quinzena (DATA_REL distintas) do cliente, em ordem crescente."""
    return sorted(pd.Timestamp(x) for x in df.loc[df["STATUS"] != "EXCLUIDO", "DATA_REL"].dropna().unique())


def efeito_do_marco(df: pd.DataFrame, ctx: Contexto, marco) -> dict:
    """O que ficaria 'anterior' e o que entraria na fila para um marco candidato."""
    ctx2 = Contexto(lista=ctx.lista, marco=parse_data(marco) or ctx.marco)
    d = preparar(df, ctx2)
    ant = d[d["DESTINO"] == D_ANTERIOR]
    sem = d[d["DESTINO"] == D_SEM_ETAPA]
    return {"n_ant": len(ant), "v_ant": float(ant["VALOR"].sum()), "n_fila": len(sem), "v_fila": float(sem["VALOR"].sum())}


# ---------------------------------------------------------------------------
# Fila de regularização
# ---------------------------------------------------------------------------
@dataclass
class Item:
    linha: int
    id: int
    data: pd.Timestamp
    nome: str
    referencia: str
    valor: float
    categoria: str
    sug: Optional[str] = None
    motivo: str = ""
    conf: str = "-"          # 'Alta', 'Média' ou '-'


@dataclass
class Grupo:
    chave: str
    rotulo: str
    itens: List[Item]

    @property
    def valor(self) -> float:
        return sum(i.valor for i in self.itens)


def dominantes_por_quinzena(d: pd.DataFrame) -> Dict[pd.Timestamp, Tuple[str, float, str]]:
    """
    Etapa em andamento em cada quinzena: etapa dominante da mão de obra já classificada
    (por valor); sem mão de obra classificada, das demais linhas classificadas (por quantidade).
    Devolve {data: (etapa, participação 0-1, base)}.
    """
    out = {}
    cl = d[d["DESTINO"] == D_ETAPA]
    for dt, g in cl.groupby("DATA_REL"):
        mo = g[g["CATEGORIA"] == "MO"]
        if len(mo) and mo["VALOR"].sum() > 0:
            v = mo.groupby("ETAPA_FINAL")["VALOR"].sum().sort_values(ascending=False)
            out[dt] = (v.index[0], float(v.iloc[0] / v.sum()), "MO classificada")
        else:
            v = g.groupby("ETAPA_FINAL").size().sort_values(ascending=False)
            out[dt] = (v.index[0], float(v.iloc[0] / v.sum()), "demais lançamentos")
    return out


def _item(r) -> Item:
    return Item(linha=int(r.LINHA), id=int(r.ID) if pd.notna(r.ID) else -1, data=r.DATA_REL, nome=r.NOME,
                referencia=r.REFERENCIA, valor=float(r.VALOR), categoria=r.CATEGORIA)


def fila_mao_de_obra(d: pd.DataFrame) -> List[Grupo]:
    """Mão de obra sem etapa (a partir do marco), agrupada por quinzena, com a etapa em andamento como sugestão."""
    pend = d[(d["CATEGORIA"] == "MO") & (d["DESTINO"] == D_SEM_ETAPA) & (~d["ETAPA_INVALIDA"])
             & (d["ETAPA_RAW"] == "")]
    dom = dominantes_por_quinzena(d)
    grupos = []
    for dt, g in pend.groupby("DATA_REL"):
        itens = [_item(r) for r in g.itertuples(index=False)]
        if dt in dom:
            etapa, share, base = dom[dt]
            for it in itens:
                if share >= LIMITE_DOMINANCIA:
                    it.sug, it.conf = etapa, "Alta" if base == "MO classificada" else "Média"
                    origem = "da MO já classificada na quinzena" if base == "MO classificada" else "dos demais lançamentos da quinzena"
                    it.motivo = f"{share * 100:.0f}% {origem}"
                else:
                    it.motivo = f"quinzena dividida: {etapa} {share * 100:.0f}% (escolha)"
        else:
            for it in itens:
                it.motivo = "nada classificado na quinzena"
        grupos.append(Grupo(chave=f"Q|{pd.Timestamp(dt).date()}", rotulo=fmt_data(dt), itens=itens))
    return sorted(grupos, key=lambda x: x.chave)


# Palavras-chave para sugerir etapa a partir da referência / descrição de contrato
_PALAVRAS_CHAVE = [
    (r"ESQUADRI|VIDRO|\bBOX\b|GUARDA.?CORPO", "ESQUADRIAS"),
    (r"IMPERMEAB|PENETRON|MANTA", "IMPERMEABILIZAÇÃO"),
    (r"ELETRIC|LUMINARIA|FIACAO", "INSTALAÇÕES ELÉTRICAS"),
    (r"HIDRAUL|ESGOTO|AGUA FRIA|AGUA QUENTE", "INSTALAÇÕES HIDRÁULICAS"),
    (r"\bPINTURA|\bTINTA|MASSA CORRIDA|EMASSAMENTO", "PINTURA"),
    (r"\bPISOS?\b|PORCELANATO", "PISOS"),
    (r"TELHA|\bCALHA|COBERTURA|\bRUFO", "COBERTURA"),
    (r"PAISAGISM|\bGRAMA\b|JARDIM", "PAISAGISMO"),
    (r"\bPROJETO|\bART\b|\bCREA\b|\bRRT\b|ALVARA|LICENCA|CARTORIO|PAPELARIA", ETAPA_GERAIS),
]
_PALAVRAS_CONTRATO = [
    (r"ELETRIC|LUMINARIA", "INSTALAÇÕES ELÉTRICAS"),
    (r"HIDRAUL|AGUA FRIA|ESGOTO", "INSTALAÇÕES HIDRÁULICAS"),
    (r"ACABAMENTO", "ACABAMENTOS"),
]


def carregar_contratos(arquivo) -> pd.DataFrame:
    """Contratos de medição (empreiteiros). Vazio se a aba não existir."""
    try:
        return pd.read_excel(arquivo, sheet_name="Contratos_Medicao")
    except Exception:  # noqa: BLE001
        return pd.DataFrame(columns=["ID_Contrato", "Nome_Fornecedor", "Descricao"])


def _etapa_por_regex(texto, tabela, ctx: Contexto) -> Optional[str]:
    t = normalizar(texto)
    for pat, etapa in tabela:
        if re.search(pat, t):
            return ctx.canonica(etapa)      # só sugere etapa que exista na lista do sistema
    return None


def _etapa_por_nome_literal(texto, ctx: Contexto) -> Optional[str]:
    return ctx.canonica(texto) if normalizar(texto) else None


def fila_fornecedores(d: pd.DataFrame, contratos: pd.DataFrame, ctx: Contexto) -> List[Grupo]:
    """
    Material, locação, serviço e diversos sem etapa, agrupados por fornecedor.
    Ordem das sugestões: 1) contrato de medição; 2) fornecedor já classificado
    (>= 3 lançamentos e >= 80% na mesma etapa); 3) palavra-chave / nome de etapa na referência.
    Mensalidade de motoboy (DIV): etapa em andamento na quinzena. A data não é critério para os demais.
    """
    pend = d[(d["CATEGORIA"].isin(["MAT", "LOC", "SERV", "DIV"])) & (d["DESTINO"] == D_SEM_ETAPA)
             & (~d["ETAPA_INVALIDA"]) & (d["ETAPA_RAW"] == "")]
    cl = d[d["DESTINO"] == D_ETAPA]
    hist = cl.groupby(["NOME", "ETAPA_FINAL"]).size()
    hist_nomes = set(hist.index.get_level_values(0))
    dom = dominantes_por_quinzena(d)
    ctr = {}
    for nome, g in contratos.groupby(contratos["Nome_Fornecedor"].map(normalizar)):
        ets = {_etapa_por_regex(x, _PALAVRAS_CONTRATO, ctx) for x in g["Descricao"].astype(str)}
        ets.discard(None)
        ctr[nome] = (ets.pop() if len(ets) == 1 else None, ", ".join(str(int(i)) for i in g["ID_Contrato"] if pd.notna(i)))
    grupos = []
    for nome, g in pend.groupby("NOME"):
        itens = [_item(r) for r in g.itertuples(index=False)]
        sug, motivo, conf = None, "nenhuma etapa da lista corresponde", "-"
        nn = normalizar(nome)
        if nn in ctr and ctr[nn][0]:
            sug, conf = ctr[nn][0], "Alta"
            fora = int(hist[nome][hist[nome].index != sug].sum()) if nome in hist_nomes else 0
            motivo = f"contrato de medição nº {ctr[nn][1]}" + (f"; {fora} já classificado(s) em outras etapas" if fora else "")
        elif nome in hist_nomes and hist[nome].sum() >= 3 and hist[nome].max() / hist[nome].sum() >= LIMITE_DOMINANCIA:
            sug, conf = hist[nome].idxmax(), "Alta"
            motivo = f"fornecedor já classificado ({int(hist[nome].max())}x)"
        if sug:
            for it in itens:
                it.sug, it.motivo, it.conf = sug, motivo, conf
        else:
            for it in itens:
                if it.categoria == "DIV" and re.search(r"MOTOBOY", normalizar(it.nome + " " + it.referencia)):
                    if it.data in dom and dom[it.data][1] >= LIMITE_DOMINANCIA:
                        it.sug, it.conf, it.motivo = dom[it.data][0], "Média", "mensalidade: etapa em andamento na quinzena"
                    else:
                        it.motivo = "mensalidade: quinzena sem etapa dominante"
                    continue
                e = _etapa_por_nome_literal(it.referencia, ctx) or _etapa_por_regex(it.referencia, _PALAVRAS_CHAVE, ctx)
                if e:
                    it.sug, it.conf, it.motivo = e, "Média", "palavra-chave na referência"
                else:
                    it.motivo = motivo
        grupos.append(Grupo(chave=f"F|{nome}", rotulo=nome, itens=itens))
    return sorted(grupos, key=lambda x: -x.valor)


def auditoria(d: pd.DataFrame, contratos: pd.DataFrame, ctx: Contexto) -> List[dict]:
    """Lista (sem alterar nada) de lançamentos cuja etapa gravada conflita com as regras."""
    achados = []
    for r in d.itertuples(index=False):
        if r.ETAPA_RAW and r.CATEGORIA in CATS_REGRA:
            achados.append(_achado(r, "Categoria " + r.CATEGORIA + ": no relatório vai para Custos Gerais (a etapa gravada é ignorada)"))
        elif r.ETAPA_INVALIDA:
            achados.append(_achado(r, "Etapa fora da lista de Configurações (tratada como sem etapa informada)"))
    if len(contratos):
        ctr = {}
        for nome, g in contratos.groupby(contratos["Nome_Fornecedor"].map(normalizar)):
            ets = {_etapa_por_regex(x, _PALAVRAS_CONTRATO, ctx) for x in g["Descricao"].astype(str)}
            ets.discard(None)
            if len(ets) == 1:
                ctr[nome] = ets.pop()
        x = d[(d["DESTINO"] == D_ETAPA) & (d["CATEGORIA"] == "SERV")]
        for r in x.itertuples(index=False):
            e = ctr.get(normalizar(r.NOME))
            if e and r.ETAPA_FINAL != e:
                achados.append(_achado(r, f"Etapa diverge do contrato de medição (contrato: {e})"))
    return sorted(achados, key=lambda a: (a["situacao"], a["data"] if pd.notna(a["data"]) else pd.Timestamp.min))


def _achado(r, situacao) -> dict:
    return {"linha": int(r.LINHA), "data": r.DATA_REL, "categoria": r.CATEGORIA, "nome": r.NOME,
            "referencia": r.REFERENCIA, "valor": float(r.VALOR), "etapa": r.ETAPA_RAW, "situacao": situacao}


# ---------------------------------------------------------------------------
# Gravação segura
# ---------------------------------------------------------------------------
def _pasta_backup(arquivo) -> Path:
    return Path(arquivo).parent / PASTA_BACKUP


def _editar_planilha(arquivo, editar: Callable, sufixo_backup: str,
                     verificar: Optional[Callable] = None,
                     precisa_salvar: Optional[Callable] = None) -> Tuple[object, str]:
    """
    Fluxo único de gravação:
      1) testa se o arquivo está livre; 2) copia backup; 3) edita uma cópia temporária;
      4) confere a cópia; 5) confere que ninguém alterou o original; 6) troca atômica.
    Se qualquer etapa falhar, o arquivo original permanece intacto.
    """
    arquivo = Path(arquivo)
    try:
        with open(arquivo, "r+b"):
            pass
    except (PermissionError, OSError) as e:
        raise ArquivoEmUsoError(f"A planilha está aberta em outro programa ou bloqueada:\n{arquivo}") from e

    mtime0 = arquivo.stat().st_mtime_ns
    pasta = _pasta_backup(arquivo)
    pasta.mkdir(parents=True, exist_ok=True)
    backup = pasta / f"{arquivo.stem}_{sufixo_backup}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{arquivo.suffix}"
    shutil.copy2(arquivo, backup)

    wb = load_workbook(arquivo)
    tmp = arquivo.with_name(f"~tmp_etapas_{arquivo.name}")
    try:
        resultado = editar(wb)
        if precisa_salvar is not None and not precisa_salvar(resultado):
            backup.unlink(missing_ok=True)          # nada a gravar: descarta o backup
            return resultado, ""
        wb.save(tmp)
    finally:
        wb.close()
    try:
        if verificar is not None:
            verificar(tmp, resultado)
        else:
            load_workbook(tmp, read_only=True).close()
        if arquivo.stat().st_mtime_ns != mtime0:
            raise PlanilhaAlteradaError("A planilha foi alterada por outro usuário durante a operação. Nada foi gravado.")
        try:
            os.replace(tmp, arquivo)
        except PermissionError as e:
            raise ArquivoEmUsoError(f"A planilha foi aberta em outro programa durante a gravação:\n{arquivo}") from e
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
    return resultado, str(backup)


@dataclass
class ResultadoGravacao:
    gravados: List[dict] = field(default_factory=list)
    pulados: List[Tuple[int, str]] = field(default_factory=list)
    backup: str = ""


def _vazio(v) -> bool:
    return v is None or str(v).strip() == ""


def gravar_etapas(arquivo, alteracoes: List[dict], ctx: Optional[Contexto] = None) -> ResultadoGravacao:
    """
    Grava ETAPA_OBRA (coluna Q) e HISTORICO_ALTERACAO (coluna P) das linhas informadas.

    alteracoes: [{'linha': nº da linha no Excel, 'id': ID_LANCAMENTO, 'valor': VALOR, 'etapa': texto}]

    Cada linha é conferida imediatamente antes de gravar (ID, valor, status ATIVO e coluna Q vazia);
    se algo mudou, a linha é pulada e informada. Nunca sobrescreve etapa já preenchida, não ordena,
    não insere e não apaga linhas, e não toca em outras colunas ou abas.
    """
    ctx = ctx or criar_contexto(arquivo)
    resultado = ResultadoGravacao()
    timestamp = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    def editar(wb):
        if ABA_DADOS not in wb.sheetnames:
            raise ValueError(f"A planilha não possui a aba '{ABA_DADOS}'.")
        ws = wb[ABA_DADOS]
        cols = _mapa_colunas(ws)
        c_val, c_st, c_id = cols["VALOR"], cols["STATUS"], cols["ID_LANCAMENTO"]
        if (c_id, cols["HISTORICO_ALTERACAO"], cols["ETAPA_OBRA"]) != (COL_ID, COL_HISTORICO, COL_ETAPA):
            raise ValueError("Layout inesperado da aba Dados (colunas ID_LANCAMENTO/HISTORICO_ALTERACAO/ETAPA_OBRA "
                             "fora das posições O/P/Q). Nada foi gravado.")
        cab = ws.cell(row=1, column=COL_ETAPA).value
        if _vazio(cab):
            ws.cell(row=1, column=COL_ETAPA, value="ETAPA_OBRA")
        elif normalizar(cab) != "ETAPA_OBRA":
            raise ValueError(f"A coluna Q da aba Dados tem o cabeçalho '{cab}' (esperado ETAPA_OBRA). Nada foi gravado.")
        for alt in alteracoes:
            r, etapa = int(alt["linha"]), str(alt["etapa"]).strip()
            canon = ctx.canonica(etapa)
            if not canon:
                resultado.pulados.append((r, f"etapa '{etapa}' não está na lista de etapas")); continue
            if r < 2 or r > ws.max_row:
                resultado.pulados.append((r, "linha inexistente")); continue
            id_atual = ws.cell(row=r, column=c_id).value
            try:
                mesmo_id = int(id_atual) == int(alt["id"])
            except (TypeError, ValueError):
                mesmo_id = False
            if not mesmo_id:
                resultado.pulados.append((r, "o ID da linha mudou (planilha alterada por outro usuário?)")); continue
            if normalizar(ws.cell(row=r, column=c_st).value) == "EXCLUIDO":
                resultado.pulados.append((r, "lançamento excluído")); continue
            try:
                if abs(float(ws.cell(row=r, column=c_val).value) - float(alt["valor"])) > 0.005:
                    resultado.pulados.append((r, "o valor da linha mudou")); continue
            except (TypeError, ValueError):
                resultado.pulados.append((r, "valor ilegível")); continue
            if not _vazio(ws.cell(row=r, column=COL_ETAPA).value):
                resultado.pulados.append((r, "a linha já possui etapa")); continue
            ws.cell(row=r, column=COL_ETAPA, value=canon)
            atual = ws.cell(row=r, column=COL_HISTORICO).value
            entrada = f"EDIÇÃO: ETAPA_OBRA:  → {canon[:30] + '...' if len(canon) > 30 else canon} - {timestamp}"
            novo = entrada
            if not _vazio(atual):
                partes = str(atual).split(" | ")
                if len(partes) >= 5:
                    partes = partes[-4:]
                novo = " | ".join(partes) + " | " + entrada
            ws.cell(row=r, column=COL_HISTORICO, value=novo)
            resultado.gravados.append({"linha": r, "id": int(alt["id"]), "etapa": canon})
        return resultado

    def verificar(tmp, res):
        wb2 = load_workbook(tmp, read_only=True)
        try:
            ws2 = wb2[ABA_DADOS]
            for g in res.gravados:
                if normalizar(ws2.cell(row=g["linha"], column=COL_ETAPA).value) != normalizar(g["etapa"]):
                    raise ValueError(f"Conferência pós-gravação falhou na linha {g['linha']}. Nada foi gravado.")
        finally:
            wb2.close()

    if not alteracoes:
        return resultado
    _, resultado.backup = _editar_planilha(arquivo, editar, "ANTES_ETAPAS", verificar,
                                        precisa_salvar=lambda r: bool(r.gravados))
    return resultado
