# -*- coding: utf-8 -*-
"""
caixa_entrada.py - A planilha Documentos_Recebidos.xlsx.

Ela fica ENTRE o Gmail e as planilhas dos clientes. Nada é gravado direto na
planilha de um cliente: o robô só sugere; quem confirma é você, pela tela de
Entrada de Dados (botão "📧 E-mail"), passando pelas validações que o sistema
já tem (duplicidade, cliente errado, ID sequencial, backup).

Abas:
  Lancamentos  -> uma linha por conta a pagar sugerida (parcela/boleto)
  Documentos   -> um registro por documento lido (evita ler o mesmo 2 vezes)
  Mensagens    -> e-mails já processados (evita baixar de novo)
  Listas       -> apoio às listas suspensas (clientes, status)
  LEIA-ME      -> legenda
"""

VERSAO = "2026-09-26.4"   # igual em todos os arquivos do módulo
import time
from datetime import datetime, date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from . import config_email as cfg
from src.salvamento_seguro import salvar_planilha_segura

COLUNAS_LANC = [
    "ID_SUGESTAO", "STATUS", "CLIENTE", "IDENTIFICACAO", "MOTIVO_IDENTIFICACAO",
    "TP_DESP", "CNPJ_CPF", "NOME", "REFERENCIA", "NF", "VALOR", "DT_VENCTO",
    "CATEGORIA", "DADOS_BANCARIOS", "ALERTAS", "TIPO_DOC", "CHAVE_DOC",
    "LINHA_DIGITAVEL", "ARQUIVOS", "EMAIL_DATA", "EMAIL_REMETENTE",
    "EMAIL_ASSUNTO", "MSG_ID", "DATA_CAPTURA", "DATA_IMPORTACAO",
]
COLUNAS_DOC = ["CHAVE", "TIPO", "NUMERO", "EMIT_DOC", "VALOR", "MSG_ID", "ARQUIVO", "DATA_CAPTURA"]
COLUNAS_MSG = ["MSG_ID", "DATA_EMAIL", "REMETENTE", "ASSUNTO", "RESULTADO", "DATA_PROCESSAMENTO"]

STATUS_VALIDOS = ["PENDENTE", "IMPORTADO", "JA_LANCADO", "IGNORADO", "CANCELADA"]

LARGURAS = {
    "ID_SUGESTAO": 8, "STATUS": 12, "CLIENTE": 34, "IDENTIFICACAO": 17,
    "MOTIVO_IDENTIFICACAO": 34, "TP_DESP": 8, "CNPJ_CPF": 17, "NOME": 32,
    "REFERENCIA": 34, "NF": 10, "VALOR": 12, "DT_VENCTO": 12, "CATEGORIA": 10,
    "DADOS_BANCARIOS": 40, "ALERTAS": 50, "TIPO_DOC": 10, "CHAVE_DOC": 46,
    "LINHA_DIGITAVEL": 50, "ARQUIVOS": 50, "EMAIL_DATA": 17,
    "EMAIL_REMETENTE": 30, "EMAIL_ASSUNTO": 40, "MSG_ID": 18,
    "DATA_CAPTURA": 17, "DATA_IMPORTACAO": 17,
}


class ArquivoEmUso(Exception):
    pass


class CaixaEntrada:
    def __init__(self, arquivo=None):
        self.arquivo = arquivo or cfg.ARQUIVO_CAIXA
        self.lancamentos, self.documentos, self.mensagens = [], [], []
        if self.arquivo.exists():
            self._carregar()

    # ------------------------------------------------------------------ leitura
    def _carregar(self):
        wb = load_workbook(self.arquivo)  # não read_only: preserva edições manuais
        self.lancamentos = self._ler_aba(wb, "Lancamentos")
        self.documentos = self._ler_aba(wb, "Documentos")
        self.mensagens = self._ler_aba(wb, "Mensagens")
        wb.close()

    @staticmethod
    def _ler_aba(wb, nome):
        if nome not in wb.sheetnames:
            return []
        linhas = wb[nome].iter_rows(values_only=True)
        try:
            cab = [str(c) if c is not None else "" for c in next(linhas)]
        except StopIteration:
            return []
        return [dict(zip(cab, l)) for l in linhas if any(v not in (None, "") for v in l)]

    # --------------------------------------------------------------- consultas
    def mensagens_processadas(self):
        return {str(m["MSG_ID"]) for m in self.mensagens if m.get("MSG_ID")}

    def chaves_documentos(self):
        return {str(d["CHAVE"]) for d in self.documentos if d.get("CHAVE")}

    def proximo_id(self):
        ids = [int(l["ID_SUGESTAO"]) for l in self.lancamentos
               if str(l.get("ID_SUGESTAO") or "").isdigit()]
        return (max(ids) if ids else 0) + 1

    def pendentes(self, cliente=None, incluir_sem_cliente=False):
        saida = []
        for l in self.lancamentos:
            if str(l.get("STATUS") or "").upper() != "PENDENTE":
                continue
            cli = str(l.get("CLIENTE") or "").strip()
            if cliente is None or cli == cliente or (incluir_sem_cliente and not cli):
                saida.append(l)
        return saida

    # ---------------------------------------------------------------- inclusão
    def registrar_mensagem(self, msg_id, data_email, remetente, assunto, resultado):
        self.mensagens.append({
            "MSG_ID": msg_id, "DATA_EMAIL": data_email, "REMETENTE": remetente,
            "ASSUNTO": assunto, "RESULTADO": resultado,
            "DATA_PROCESSAMENTO": datetime.now(),
        })

    def registrar_documento(self, chave, tipo, numero, emit_doc, valor, msg_id, arquivo):
        self.documentos.append({
            "CHAVE": chave, "TIPO": tipo, "NUMERO": numero, "EMIT_DOC": emit_doc,
            "VALOR": valor, "MSG_ID": msg_id, "ARQUIVO": arquivo,
            "DATA_CAPTURA": datetime.now(),
        })

    def adicionar_lancamento(self, dados: dict):
        dados = dict(dados)
        dados["ID_SUGESTAO"] = self.proximo_id()
        dados.setdefault("STATUS", "PENDENTE")
        dados.setdefault("DATA_CAPTURA", datetime.now())
        self.lancamentos.append(dados)
        return dados["ID_SUGESTAO"]

    def buscar_lancamento_sem_boleto(self, emit_doc, valor, vencimento):
        """Para casar um boleto que chegou DEPOIS da nota, em outro e-mail."""
        from .roteador import _chave_doc
        k = _chave_doc(emit_doc)
        for l in self.lancamentos:
            if str(l.get("STATUS") or "").upper() != "PENDENTE":
                continue
            if l.get("LINHA_DIGITAVEL") or _chave_doc(l.get("CNPJ_CPF")) != k:
                continue
            try:
                v = float(l.get("VALOR") or 0)
            except (TypeError, ValueError):
                continue
            if abs(v - (valor or 0)) > cfg.TOLERANCIA_VALOR:
                continue
            venc = l.get("DT_VENCTO")
            if isinstance(venc, datetime):
                venc = venc.date()
            if venc and vencimento and venc != vencimento:
                continue
            return l
        return None

    def unir_boletos_e_notas(self):
        """Boleto e nota do mesmo fornecedor que chegaram em e-mails DIFERENTES
        (em qualquer ordem) viram UMA sugestão só: a da nota, que recebe a
        linha digitável; a do boleto fica IGNORADO com "UNIDO ao ID n".
        Só une quando o par é inequívoco: mesmo CNPJ, mesmo valor, mesmo
        vencimento (ou nota sem vencimento), mesmo cliente, um único candidato."""
        from .roteador import _chave_doc
        from datetime import datetime as _dt

        def dia(v):
            return v.date() if isinstance(v, _dt) else v

        def valor(v):
            try:
                return round(float(v), 2)
            except (TypeError, ValueError):
                return None

        unidos = 0
        pend = [l for l in self.lancamentos if str(l.get("STATUS") or "").upper() == "PENDENTE"]
        boletos = [l for l in pend if l.get("LINHA_DIGITAVEL") and str(l.get("TIPO_DOC")) == "BOLETO"
                   and not l.get("NF")]
        notas = [l for l in pend if not l.get("LINHA_DIGITAVEL") and str(l.get("TIPO_DOC")) in ("NFE", "NFSE")]
        usadas = set()
        for b in boletos:
            kb, vb, db = _chave_doc(b.get("CNPJ_CPF")), valor(b.get("VALOR")), dia(b.get("DT_VENCTO"))
            if not kb or vb is None:
                continue
            cands = []
            for n in notas:
                if id(n) in usadas or _chave_doc(n.get("CNPJ_CPF")) != kb:
                    continue
                vn, dn = valor(n.get("VALOR")), dia(n.get("DT_VENCTO"))
                if vn is None or abs(vn - vb) > cfg.TOLERANCIA_VALOR:
                    continue
                if dn and db and dn != db:
                    continue
                cb, cn = str(b.get("CLIENTE") or ""), str(n.get("CLIENTE") or "")
                if cb and cn and cb != cn:
                    continue
                cands.append(n)
            if len(cands) != 1:
                continue
            n = cands[0]
            usadas.add(id(n))
            n["LINHA_DIGITAVEL"] = b["LINHA_DIGITAVEL"]
            n["DADOS_BANCARIOS"] = b.get("DADOS_BANCARIOS") or n.get("DADOS_BANCARIOS")
            if not n.get("DT_VENCTO") and db:
                n["DT_VENCTO"] = db
                n["ALERTAS"] = str(n.get("ALERTAS") or "").replace("Sem vencimento - preencha DT_VENCTO", "")
            if not n.get("CLIENTE") and b.get("CLIENTE"):
                n["CLIENTE"], n["IDENTIFICACAO"] = b["CLIENTE"], b.get("IDENTIFICACAO")
            extras = [a.strip() for a in str(b.get("ALERTAS") or "").split("|") if a.strip().startswith(("ATENÇÃO", "Vencimento do código"))]
            n["ALERTAS"] = " | ".join(x.strip() for x in str(n.get("ALERTAS") or "").split("|") + extras if x.strip())
            n["ARQUIVOS"] = " ; ".join(x for x in [n.get("ARQUIVOS"), b.get("ARQUIVOS")] if x)
            b["STATUS"] = "IGNORADO"
            b["ALERTAS"] = f"UNIDO ao ID {n['ID_SUGESTAO']} (nota {n.get('NF')}) | {b.get('ALERTAS') or ''}".strip(" |")
            unidos += 1
        return unidos

    # ---------------------------------------------------------------- gravação
    def salvar(self, nomes_clientes=None, tentativas=3):
        wb = Workbook()
        wb.remove(wb.active)
        self._escrever_aba(wb, "Lancamentos", COLUNAS_LANC, self.lancamentos)
        self._escrever_aba(wb, "Documentos", COLUNAS_DOC, self.documentos)
        self._escrever_aba(wb, "Mensagens", COLUNAS_MSG, self.mensagens)
        self._escrever_listas(wb, nomes_clientes or [])
        self._escrever_leiame(wb)

        self.arquivo.parent.mkdir(parents=True, exist_ok=True)
        for i in range(tentativas):
            try:
                salvar_planilha_segura(wb, self.arquivo)
                return
            except PermissionError:
                if i == tentativas - 1:
                    raise ArquivoEmUso(
                        f"Não foi possível salvar {self.arquivo.name}: o arquivo está aberto "
                        f"(Excel ou Google Drive sincronizando). Feche-o e rode de novo. "
                        f"Nada foi perdido: os e-mails serão lidos novamente.")
                time.sleep(3)

    @staticmethod
    def _escrever_aba(wb, nome, colunas, linhas):
        ws = wb.create_sheet(nome)
        ws.append(colunas)
        for l in linhas:
            ws.append([l.get(c) for c in colunas])
        cab_fonte = Font(name="Arial", bold=True, color="FFFFFF")
        cab_fundo = PatternFill("solid", fgColor="1F4E78")
        for i, c in enumerate(colunas, start=1):
            cel = ws.cell(row=1, column=i)
            cel.font, cel.fill = cab_fonte, cab_fundo
            cel.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = LARGURAS.get(c, 18)
        fonte = Font(name="Arial", size=10)
        for linha in ws.iter_rows(min_row=2):
            for cel in linha:
                cel.font = fonte
                if isinstance(cel.value, datetime):
                    cel.number_format = "dd/mm/yyyy hh:mm"
                elif isinstance(cel.value, date):
                    cel.number_format = "dd/mm/yyyy"
                elif isinstance(cel.value, float):
                    cel.number_format = "#,##0.00"
        # colunas que são SEMPRE texto (evita perder zeros à esquerda e virar notação científica)
        for c in ("CNPJ_CPF", "CHAVE_DOC", "LINHA_DIGITAVEL", "CHAVE", "EMIT_DOC", "NF", "NUMERO"):
            if c in colunas:
                col = colunas.index(c) + 1
                for r in range(2, ws.max_row + 1):
                    cel = ws.cell(row=r, column=col)
                    if cel.value is not None:
                        cel.value = str(cel.value)
                    cel.number_format = "@"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    @staticmethod
    def _escrever_listas(wb, nomes_clientes):
        ws = wb.create_sheet("Listas")
        ws["A1"], ws["B1"] = "CLIENTES", "STATUS"
        for i, n in enumerate(nomes_clientes, start=2):
            ws.cell(row=i, column=1, value=n)
        for i, s in enumerate(STATUS_VALIDOS, start=2):
            ws.cell(row=i, column=2, value=s)
        ws.column_dimensions["A"].width = 50
        ws.sheet_state = "hidden"

        lanc = wb["Lancamentos"]
        fim = 5000
        col_status = get_column_letter(COLUNAS_LANC.index("STATUS") + 1)
        col_cliente = get_column_letter(COLUNAS_LANC.index("CLIENTE") + 1)
        dv_s = DataValidation(type="list", formula1=f"=Listas!$B$2:$B${len(STATUS_VALIDOS) + 1}",
                              allow_blank=False)
        dv_s.add(f"{col_status}2:{col_status}{fim}")
        lanc.add_data_validation(dv_s)
        if nomes_clientes:
            dv_c = DataValidation(type="list", formula1=f"=Listas!$A$2:$A${len(nomes_clientes) + 1}",
                                  allow_blank=True)
            dv_c.add(f"{col_cliente}2:{col_cliente}{fim}")
            lanc.add_data_validation(dv_c)

    @staticmethod
    def _escrever_leiame(wb):
        ws = wb.create_sheet("LEIA-ME")
        textos = [
            "COMO USAR ESTA PLANILHA",
            "Ela é preenchida automaticamente pelo coletor do Gmail. Você só mexe na aba Lancamentos.",
            "",
            "CLIENTE: se estiver vazio ou errado, escolha o cliente na lista suspensa.",
            "STATUS: PENDENTE = aguardando importação | IMPORTADO = já está na planilha do cliente |",
            "        JA_LANCADO = já existia lançamento igual antes do e-mail | IGNORADO = não importar |",
            "        CANCELADA = NF-e cancelada na SEFAZ (evento recebido por e-mail).",
            "IDENTIFICACAO: CONFIRMADO = CPF do destinatário/pagador bate com o cliente;",
            "        SUGERIDO = achado pelo nome ou CNO no texto (confira!); AMBIGUO / NAO_IDENTIFICADO = escolha o cliente.",
            "ALERTAS: leia antes de importar (fornecedor não cadastrado, sem vencimento, NF sem protocolo etc.).",
            "",
            "Pode editar TP_DESP, NOME, REFERENCIA, NF, VALOR, DT_VENCTO, CATEGORIA, DADOS_BANCARIOS.",
            "Não edite ID_SUGESTAO, MSG_ID nem as abas Documentos/Mensagens.",
            "Feche esta planilha antes de rodar o coletor ou de importar pelo sistema.",
        ]
        for i, t in enumerate(textos, start=1):
            c = ws.cell(row=i, column=1, value=t)
            c.font = Font(name="Arial", bold=(i == 1), size=12 if i == 1 else 10)
        ws.column_dimensions["A"].width = 120
        wb.move_sheet("LEIA-ME", offset=-4)
        wb.active = wb.sheetnames.index("Lancamentos")
